"""
CiteCheck AI — Unit Tests for Microsoft Word .docx Parser.
Tests valid docx extraction (paragraphs, tables, headers, footers)
and verifies rejection of .doc, corrupt files, empty files, and size limits.
"""

import io
import pytest
from fastapi.testclient import TestClient
import docx

from app.main import app

client = TestClient(app)


def create_sample_docx(
    paragraphs: list[str] = None,
    table_data: list[list[str]] = None,
    header_text: str = None,
    footer_text: str = None,
) -> bytes:
    doc = docx.Document()
    if header_text:
        doc.sections[0].header.paragraphs[0].text = header_text
    if footer_text:
        doc.sections[0].footer.paragraphs[0].text = footer_text
    for p in (paragraphs or ["Sample paragraph text for testing"]):
        doc.add_paragraph(p)
    if table_data:
        tbl = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        for r_idx, row in enumerate(table_data):
            for c_idx, val in enumerate(row):
                tbl.cell(r_idx, c_idx).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class TestDocxParser:
    """Test suite for .docx extraction endpoint."""

    def test_valid_docx_with_tables_and_headers(self):
        content = create_sample_docx(
            paragraphs=["Paragraph 1", "Paragraph 2 with claims"],
            table_data=[["Header A", "Header B"], ["Row 1 Val", "Row 2 Val"]],
            header_text="Document Header Confirmed",
            footer_text="Confidential Footer",
        )
        response = client.post(
            "/api/drafts/parse-docx",
            files={"file": ("test_doc.docx", content, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert response.status_code == 200
        data = response.json()
        assert data["filename"] == "test_doc.docx"
        assert "Paragraph 1" in data["text"]
        assert "Header A | Header B" in data["text"]
        assert "Document Header Confirmed" in data["text"]
        assert "Confidential Footer" in data["text"]
        assert data["table_count"] == 1
        assert data["paragraph_count"] >= 2

    def test_reject_legacy_doc_format(self):
        response = client.post(
            "/api/drafts/parse-docx",
            files={"file": ("legacy_document.doc", b"fake binary doc content", "application/msword")},
        )
        assert response.status_code == 400
        assert "Legacy .doc format is not supported" in response.json()["detail"]

    def test_reject_unsupported_extension(self):
        response = client.post(
            "/api/drafts/parse-docx",
            files={"file": ("image.png", b"fake png bytes", "image/png")},
        )
        assert response.status_code == 400
        assert "Unsupported file format" in response.json()["detail"]

    def test_reject_empty_file(self):
        response = client.post(
            "/api/drafts/parse-docx",
            files={"file": ("empty.docx", b"", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert response.status_code == 400
        assert "empty (0 bytes)" in response.json()["detail"]

    def test_reject_corrupt_or_renamed_docx(self):
        # A text file renamed to .docx is not a valid zip archive
        response = client.post(
            "/api/drafts/parse-docx",
            files={"file": ("fake.docx", b"This is just plain text renamed to docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert response.status_code == 400
        assert "Corrupted or invalid document archive" in response.json()["detail"]

    def test_reject_oversized_file(self):
        # > 5 MB
        oversized = b"P" * (5 * 1024 * 1024 + 10)
        response = client.post(
            "/api/drafts/parse-docx",
            files={"file": ("big.docx", oversized, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert response.status_code == 413
        assert "exceeds the 5 MB limit" in response.json()["detail"]

